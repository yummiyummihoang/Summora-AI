"""Document import, normalization, budgets and summary execution."""

import io
import json
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace

from summora.documents import import_document
from summora.tasks import (
    execute_task,
    length_budget,
    normalize_payload,
    render_template,
    run_tasks,
    word_count,
)

from tests.helpers import FakeClient, payload, settings


class UnitTests(unittest.TestCase):
    def test_length_rounding_and_validation(self):
        self.assertEqual(word_count(" nghiên\tcứu\n mới  "), 3)
        self.assertEqual(
            length_budget("x " * 11, {"mode": "percent", "value": 20})["target_words"], 3
        )
        self.assertEqual(
            length_budget("x " * 1000, {"mode": "percent", "value": 20})["target_words"], 200
        )
        self.assertEqual(
            length_budget("x " * 500, {"mode": "percent", "value": 20})["target_words"], 100
        )
        self.assertEqual(
            length_budget("x " * 200, {"mode": "words", "value": 100})["min_words"], 90
        )
        for bad in [0, 101, True, float("nan"), "20"]:
            with self.assertRaises(ValueError):
                length_budget("x " * 10, {"mode": "percent", "value": bad})
        for bad in [0, 11, 1.5, True]:
            with self.assertRaises(ValueError):
                length_budget("x " * 10, {"mode": "words", "value": bad})

    def test_template_preserves_json_and_source(self):
        budget = length_budget("x " * 100, {"mode": "percent", "value": 20})
        source = "Nguồn {N} {document} {target_words}"
        template = '{"example": "test"} {N} {target_words} {compression_percent}\n{VAN_BAN_NGUON}'
        rendered = render_template(template, source, budget, {"N": 3})
        self.assertIn('{"example": "test"} 3 20 20', rendered)
        self.assertTrue(rendered.endswith(source))
        with self.assertRaises(ValueError):
            render_template("{N}", source, budget, {})

    def test_normalization_snapshot_and_overrides(self):
        body = payload()
        body["documents"][1]["overrides"] = settings("deepseek", "model-b")
        body["documents"][1]["overrides"]["length"] = {"mode": "words", "value": 15}
        docs, tasks = normalize_payload(body)
        self.assertEqual(
            [(t["provider"], t["budget"]["target_words"]) for t in tasks],
            [("openai", 20), ("deepseek", 15)],
        )
        body["settings"]["prompts"][0]["text"] = "changed"
        self.assertNotEqual(docs[0]["settings"]["prompts"][0]["text"], "changed")
        bad = payload()
        bad["documents"][1]["id"] = "first"
        with self.assertRaises(ValueError):
            normalize_payload(bad)
        bad = payload()
        bad["documents"][0]["text"] = " "
        with self.assertRaises(ValueError):
            normalize_payload(bad)

    def test_selected_prompt_task_count_and_missing_documents(self):
        body = payload()
        body["settings"]["prompts"].append({"name": "Hai", "text": "Tóm tắt"})
        self.assertEqual(len(normalize_payload(body)[1]), 4)
        with self.assertRaises(ValueError):
            normalize_payload({"settings": settings()})

    def test_import_text_csv_json_docx_pdf(self):
        self.assertEqual(
            import_document("\ufeffVăn bản thử".encode(), "test.txt")[0]["text"], "Văn bản thử"
        )
        self.assertEqual(
            len(
                import_document(
                    b'title,body\nA,"line 1\nline 2"\nB,other', "test.csv", field="body"
                )
            ),
            2,
        )
        self.assertEqual(
            import_document(b'{"prompts":[{"name":"p","text":"test"}]}', "test.json", "prompts")[0][
                "name"
            ],
            "p",
        )
        from docx import Document

        doc = Document()
        doc.add_paragraph("Đoạn đầu")
        doc.add_table(rows=1, cols=2).cell(0, 0).text = "Ô bảng"
        doc.add_paragraph("Đoạn cuối")
        stream = io.BytesIO()
        doc.save(stream)
        text = import_document(stream.getvalue(), "test.docx")[0]["text"]
        self.assertEqual(text.count("Đoạn đầu"), 1)
        self.assertLess(text.index("Ô bảng"), text.index("Đoạn cuối"))
        from pypdf import PdfWriter
        from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

        writer = PdfWriter()
        page = writer.add_blank_page(width=200, height=200)
        font = DictionaryObject(
            {
                NameObject("/Type"): NameObject("/Font"),
                NameObject("/Subtype"): NameObject("/Type1"),
                NameObject("/BaseFont"): NameObject("/Helvetica"),
            }
        )
        page[NameObject("/Resources")] = DictionaryObject(
            {NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})}
        )
        content = DecodedStreamObject()
        content.set_data(b"BT /F1 12 Tf 20 100 Td (PDF text test) Tj ET")
        page[NameObject("/Contents")] = writer._add_object(content)
        stream = io.BytesIO()
        writer.write(stream)
        self.assertIn("PDF text test", import_document(stream.getvalue(), "test.pdf")[0]["text"])
        blank = PdfWriter()
        blank.add_blank_page(width=100, height=100)
        stream = io.BytesIO()
        blank.write(stream)
        with self.assertRaisesRegex(ValueError, "OCR"):
            import_document(stream.getvalue(), "scan.pdf")
        for content, name in [
            (b"", "empty.txt"),
            (b"broken", "bad.docx"),
            (b"bad", "bad.exe"),
            (b"\xff", "bad.txt"),
        ]:
            with self.assertRaises(ValueError):
                import_document(content, name)

    def test_provider_failure_continues_other_provider(self):
        body = payload()
        body["settings"]["prompts"].append({"name": "Hai", "text": "Tóm tắt"})
        body["documents"][1]["overrides"] = settings("deepseek", "model-b")
        docs, tasks = normalize_payload(body)
        clients = {"openai": FakeClient(fail=401), "deepseek": FakeClient()}
        with tempfile.TemporaryDirectory() as folder:
            result = run_tasks(
                docs, tasks, clients, Path(folder) / "result.json", lambda *_: None, lambda: False
            )
            self.assertEqual(result["run"]["status"], "completed_with_errors")
            self.assertEqual(len(clients["openai"].calls), 1)
            self.assertEqual(result["data"][1]["generated_summaries"][0]["status"], "success")
            self.assertNotIn("must not expose", json.dumps(result))

    def test_incomplete_and_length_mismatch_keep_content(self):
        _, tasks = normalize_payload(payload())
        client = FakeClient()

        def incomplete(**kwargs):
            return SimpleNamespace(
                choices=[
                    SimpleNamespace(
                        message=SimpleNamespace(content="Quá ngắn"), finish_reason="length"
                    )
                ],
                usage=None,
            )

        client.create = incomplete
        entry = execute_task(tasks[0], client)
        self.assertEqual(entry["status"], "incomplete")
        self.assertEqual(entry["summary"], "Quá ngắn")
        self.assertFalse(entry["length_ok"])

    def test_import_keeps_reference_summary_in_output(self):
        rows = import_document(
            b'[{"id":"source-1","document":"one two three four five","summary":"human reference"}]',
            "data.json",
        )
        body = {
            "settings": settings(),
            "documents": [{**rows[0], "id": "internal-1", "revision": 1}],
        }
        docs, tasks = normalize_payload(body)
        with tempfile.TemporaryDirectory() as folder:
            result = run_tasks(
                docs,
                tasks,
                {"openai": FakeClient()},
                Path(folder) / "result.json",
                lambda *_: None,
                lambda: False,
            )
        self.assertEqual(result["data"][0]["summary"], "human reference")
        self.assertEqual(result["data"][0]["source_id"], "source-1")
